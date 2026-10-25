
import streamlit as st
import pandas as pd

# --------------------------------------------------
# PAGE CONFIGURATION
# --------------------------------------------------
st.set_page_config(
    page_title="SkillBridge AI",
    page_icon="🚀",
    layout="wide"
)

# --------------------------------------------------
# APPLICATION HEADER
# --------------------------------------------------
st.title("🚀 SkillBridge AI")
st.subheader("Bridging the Gap Between Education and Employment")

st.write(
    "Discover entry-level career paths, compare your skills "
    "with job requirements, and identify areas for improvement."
)

st.info(
    "SkillBridge AI helps graduates explore career options, "
    "identify skills gaps, and plan their professional development."
)

st.divider()

# --------------------------------------------------
# USER PROFILE
# --------------------------------------------------
st.header("👤 Create Your Career Profile")

name = st.text_input("Your name")

qualification = st.selectbox(
    "Your highest qualification",
    [
        "Matric",
        "Certificate",
        "Diploma",
        "Bachelor's Degree",
        "Honours Degree",
        "Master's Degree",
        "Other"
    ]
)

career = st.selectbox(
    "Which career interests you?",
    [
        "Data Analyst",
        "Data Scientist",
        "Software Developer",
        "Business Intelligence Analyst"
    ]
)

available_skills = [
    "Python",
    "SQL",
    "Excel",
    "Power BI",
    "Machine Learning",
    "Git",
    "Communication",
    "Data Analysis"
]

selected_skills = st.multiselect(
    "Select the skills you currently have",
    available_skills
)

# --------------------------------------------------
# ILLUSTRATIVE JOB DATA
# --------------------------------------------------
jobs = [
    {
        "Job Title": "Junior Data Analyst",
        "Career": "Data Analyst",
        "Required Skills": ["Python", "SQL", "Excel", "Power BI"]
    },
    {
        "Job Title": "Junior Data Scientist",
        "Career": "Data Scientist",
        "Required Skills": [
            "Python", "SQL", "Machine Learning", "Data Analysis"
        ]
    },
    {
        "Job Title": "Junior Software Developer",
        "Career": "Software Developer",
        "Required Skills": ["Python", "Git", "SQL"]
    },
    {
        "Job Title": "Business Intelligence Analyst",
        "Career": "Business Intelligence Analyst",
        "Required Skills": [
            "SQL", "Excel", "Power BI", "Data Analysis"
        ]
    }
]

# --------------------------------------------------
# LEARNING RECOMMENDATIONS
# --------------------------------------------------
learning_resources = {
    "Python": (
        "Practise Python programming and build a small project."
    ),
    "SQL": (
        "Learn SELECT, WHERE, JOIN, GROUP BY and aggregate functions."
    ),
    "Excel": (
        "Practise formulas, pivot tables and data cleaning."
    ),
    "Power BI": (
        "Build a dashboard using a public dataset."
    ),
    "Machine Learning": (
        "Study model training, testing and evaluation."
    ),
    "Git": (
        "Practise version control using Git and GitHub."
    ),
    "Communication": (
        "Practise presenting your findings clearly."
    ),
    "Data Analysis": (
        "Complete an exploratory data analysis project."
    )
}

# --------------------------------------------------
# MATCHING ENGINE
# --------------------------------------------------
if st.button("🔍 Find My Job Matches", type="primary"):

    if not name.strip():
        st.warning("Please enter your name.")

    elif not selected_skills:
        st.warning("Please select at least one skill.")

    else:
        user_skills = set(selected_skills)

        # Filter roles according to the selected career.
        relevant_jobs = [
            job for job in jobs
            if job["Career"] == career
        ]

        results = []

        for job in relevant_jobs:

            required_skills = set(job["Required Skills"])

            matched_skills = required_skills.intersection(user_skills)
            missing_skills = required_skills.difference(user_skills)

            score = (
                len(matched_skills) / len(required_skills)
            ) * 100

            results.append({
                "Job Title": job["Job Title"],
                "Match Score (%)": round(score, 1),
                "Matched Skills": (
                    ", ".join(sorted(matched_skills))
                    if matched_skills else "None"
                ),
                "Skills to Develop": (
                    ", ".join(sorted(missing_skills))
                    if missing_skills else "None"
                )
            })

        results_df = pd.DataFrame(results)

        results_df = results_df.sort_values(
            by="Match Score (%)",
            ascending=False
        )

        # --------------------------------------------------
        # JOB MATCH RESULTS
        # --------------------------------------------------
        st.success(f"Welcome, {name}!")

        st.subheader("📊 Your Potential Job Matches")

        col1, col2, col3 = st.columns(3)

        best_match = results_df.iloc[0]

        with col1:
            st.metric(
                "Career Selected",
                career
            )

        with col2:
            st.metric(
                "Highest Skills Match",
                f"{best_match['Match Score (%)']}%"
            )

        with col3:
            st.metric(
                "Skills Selected",
                len(user_skills)
            )

        st.dataframe(
            results_df,
            use_container_width=True,
            hide_index=True
        )

        st.caption(
            "Match scores represent the proportion of listed required "
            "skills selected by the user. They are not the probability "
            "of getting a job. The roles and requirements are illustrative "
            "examples, not verified live vacancies."
        )

        # --------------------------------------------------
        # PERSONAL LEARNING PLAN
        # --------------------------------------------------
        st.divider()
        st.subheader("💡 Your Personal Learning Plan")

        required_skills = set(
            next(
                job["Required Skills"]
                for job in jobs
                if job["Career"] == career
            )
        )

        missing_skills = required_skills - user_skills

        st.write(
            f"Your strongest skills match is **{best_match['Job Title']}**, "
            f"with a score of **{best_match['Match Score (%)']}%**."
        )

        if missing_skills:

            st.write(
                "Consider the following actions to strengthen your "
                "profile for this career:"
            )

            for skill in sorted(missing_skills):
                st.markdown(f"#### 📚 {skill}")

                st.write(
                    learning_resources.get(
                        skill,
                        "Find a relevant course and practise the skill."
                    )
                )

        else:
            st.success(
                "You selected all the skills listed for this role! "
                "Consider building a project to demonstrate your abilities."
            )

        st.info(
            "These are general learning suggestions, not a guarantee "
            "of employment or a verified assessment of an employer's "
            "requirements."
        )

# --------------------------------------------------
# FOOTER
# --------------------------------------------------

# python
# --------------------------------------------------
# JOB APPLICATION TRACKER
# --------------------------------------------------
st.divider()

st.header("📋 My Job Application Tracker")

st.write(
    "Keep track of your job applications and monitor "
    "your progress towards employment."
)

# Store applications during the current session.
if "applications" not in st.session_state:
    st.session_state.applications = []

# Form for adding an application.
with st.form("application_form", clear_on_submit=True):

    st.subheader("Add a Job Application")

    company = st.text_input("Company name")

    job_title = st.text_input("Job title")

    application_date = st.date_input(
        "Application date"
    )

    status = st.selectbox(
        "Application status",
        [
            "Applied",
            "Interview",
            "Assessment",
            "Offer Received",
            "Rejected"
        ]
    )

    submitted = st.form_submit_button(
        "➕ Add Application",
        type="primary"
    )

    if submitted:

        if not company.strip() or not job_title.strip():
            st.warning(
                "Please enter both the company name and job title."
            )

        else:
            st.session_state.applications.append({
                "Company": company.strip(),
                "Job Title": job_title.strip(),
                "Application Date": str(application_date),
                "Status": status
            })

            st.success(
                f"Application for {job_title} at {company} added!"
            )

# Display applications.
applications = st.session_state.applications

if applications:

    st.subheader("Your Applications")

    applications_df = pd.DataFrame(applications)

    total_applications = len(applications)

    interviews = sum(
        application["Status"] == "Interview"
        for application in applications
    )

    offers = sum(
        application["Status"] == "Offer Received"
        for application in applications
    )

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Total Applications",
        total_applications
    )

    col2.metric(
        "Interviews",
        interviews
    )

    col3.metric(
        "Offers Received",
        offers
    )

    st.dataframe(
        applications_df,
        use_container_width=True,
        hide_index=True
    )

    csv_data = applications_df.to_csv(index=False).encode("utf-8")

    st.download_button(
        label="⬇️ Download My Applications",
        data=csv_data,
        file_name="skillbridge_applications.csv",
        mime="text/csv"
    )

    st.caption(
        "Application records are stored in the current session. "
        "Download your records to keep a copy. They are not yet "
        "saved permanently to a database."
    )

else:

    st.info(
        "You haven't added any applications yet. "
        "Use the form above to start tracking your job search."
    )

# --------------------------------------------------
# FOOTER
# --------------------------------------------------
st.divider()

st.caption(
    "SkillBridge AI | FNB App of the Year Hackathon 2026"
)

#python
# --------------------------------------------------
# ATS CV CHECKER
# --------------------------------------------------
st.divider()

st.header("📄 ATS CV Checker")

st.write(
    "Analyse how closely your CV matches a job description. "
    "Identify relevant keywords, review important CV sections, "
    "and discover practical ways to improve your application."
)

st.info(
    "This is an independent CV-analysis tool, not an actual "
    "employer ATS. Its keyword score is an estimate and does "
    "not predict whether an employer will shortlist you."
)

# --------------------------------------------------
# CV INPUT
# --------------------------------------------------
cv_input_method = st.radio(
    "How would you like to provide your CV?",
    [
        "Upload a PDF or Word document",
        "Paste my CV text"
    ],
    horizontal=True
)

cv_text = ""

if cv_input_method == "Upload a PDF or Word document":

    uploaded_cv = st.file_uploader(
        "Upload your CV",
        type=["pdf", "docx"],
        help="Supported formats: PDF and Word (.docx)."
    )

    if uploaded_cv is not None:

        try:

            if uploaded_cv.name.lower().endswith(".pdf"):

                import fitz

                with fitz.open(
                    stream=uploaded_cv.getvalue(),
                    filetype="pdf"
                ) as pdf_document:

                    cv_text = "\n".join(
                        page.get_text()
                        for page in pdf_document
                    )

            elif uploaded_cv.name.lower().endswith(".docx"):

                from docx import Document
                from io import BytesIO

                document = Document(
                    BytesIO(uploaded_cv.getvalue())
                )

                paragraph_text = [
                    paragraph.text
                    for paragraph in document.paragraphs
                ]

                table_text = []

                for table in document.tables:
                    for row in table.rows:
                        table_text.append(
                            " ".join(
                                cell.text for cell in row.cells
                            )
                        )

                cv_text = "\n".join(
                    paragraph_text + table_text
                )

            if not cv_text.strip():

                st.warning(
                    "No readable text was extracted. "
                    "If your CV is a scanned PDF, try a text-based "
                    "PDF or paste the CV text manually."
                )

            else:

                st.success(
                    "CV text extracted successfully. "
                    f"Characters extracted: {len(cv_text)}"
                )

        except Exception:

            st.error(
                "The document could not be read. Please check that "
                "the file is valid and try another document."
            )

            cv_text = ""

else:

    cv_text = st.text_area(
        "Paste your CV text here",
        height=250,
        placeholder=(
            "Paste your CV text, including your skills, "
            "education, experience and projects..."
        )
    )

# --------------------------------------------------
# JOB DESCRIPTION
# --------------------------------------------------
st.subheader("🎯 Target Job Description")

job_description = st.text_area(
    "Paste the job advertisement or job description",
    height=220,
    placeholder=(
        "Paste the job title, responsibilities, "
        "required skills and qualifications here..."
    )
)

# --------------------------------------------------
# ANALYSE THE CV
# --------------------------------------------------
if st.button("🔍 Analyse My CV", type="primary"):

    if not cv_text.strip():

        st.warning(
            "Please upload a readable CV or paste your CV text."
        )

    elif not job_description.strip():

        st.warning(
            "Please paste the job description before analysing your CV."
        )

    else:

        import re
        from collections import Counter

        # Normalise text for basic keyword comparison.
        def normalise_text(text):
            return re.findall(
                r"\b[a-zA-Z0-9+#.]+\b",
                text.lower()
            )

        cv_words = normalise_text(cv_text)
        job_words = normalise_text(job_description)

        # Common words that add little value to keyword matching.
        stop_words = {
            "the", "and", "for", "with", "you", "your",
            "are", "our", "this", "that", "from", "have",
            "has", "will", "can", "all", "any", "not",
            "but", "who", "what", "when", "where", "how",
            "about", "into", "their", "they", "them", "its",
            "job", "role", "work", "working", "candidate",
            "candidates", "responsibilities", "requirements",
            "required", "preferred", "ability", "including",
            "experience", "skills", "skill", "knowledge"
        }

        # Extract distinct, meaningful terms from the job description.
        job_keywords = {
            word for word in job_words
            if len(word) > 1 and word not in stop_words
        }

        cv_keyword_set = set(cv_words)

        matched_keywords = sorted(
            job_keywords.intersection(cv_keyword_set)
        )

        missing_keywords = sorted(
            job_keywords.difference(cv_keyword_set)
        )

        if job_keywords:

            match_score = (
                len(matched_keywords) / len(job_keywords)
            ) * 100

        else:

            match_score = 0

        # --------------------------------------------------
        # RESULTS
        # --------------------------------------------------
        st.divider()
        st.subheader("📊 Your CV Analysis Results")

        col1, col2, col3 = st.columns(3)

        with col1:
            st.metric(
                "Keyword Match Score",
                f"{match_score:.1f}%"
            )

        with col2:
            st.metric(
                "Matched Keywords",
                len(matched_keywords)
            )

        with col3:
            st.metric(
                "Potential Missing Keywords",
                len(missing_keywords)
            )

        st.progress(int(match_score))

        st.caption(
            "This score measures exact keyword overlap after "
            "basic text normalisation and common-word filtering. "
            "It does not measure your overall suitability for a job."
        )

        # --------------------------------------------------
        # MATCHED KEYWORDS
        # --------------------------------------------------
        st.subheader("✅ Keywords Found in Your CV")

        if matched_keywords:

            st.write(", ".join(matched_keywords))

        else:

            st.write(
                "No matching keywords were detected. "
                "Check whether the CV text was extracted correctly."
            )

        # --------------------------------------------------
        # MISSING KEYWORDS
        # --------------------------------------------------
        st.subheader("📝 Keywords to Review")

        if missing_keywords:

            st.write(
                "These terms appear in the job description but "
                "were not detected in your CV text:"
            )

            st.write(", ".join(missing_keywords))

            st.warning(
                "Review these terms carefully. Add a skill or "
                "experience only if it truthfully reflects your "
                "background. Some terms may not be relevant to you."
            )

        else:

            st.success(
                "All detected job-description keywords were found "
                "in your CV text."
            )

        # --------------------------------------------------
        # BASIC CV SECTION CHECK
        # --------------------------------------------------
        st.subheader("📋 CV Structure Check")

        cv_lower = cv_text.lower()

        section_checks = {
            "Contact information": any(
                term in cv_lower
                for term in [
                    "@", "phone", "telephone", "contact"
                ]
            ),
            "Education": any(
                term in cv_lower
                for term in [
                    "education", "qualification", "degree",
                    "university", "college"
                ]
            ),
            "Experience or projects": any(
                term in cv_lower
                for term in [
                    "experience", "employment", "projects",
                    "internship"
                ]
            ),
            "Skills": "skills" in cv_lower,
        }

        for section, found in section_checks.items():

            if found:

                st.write(f"✅ {section}: possible section detected")

            else:

                st.write(
                    f"⚠️ {section}: not clearly detected; review your CV"
                )

        # --------------------------------------------------
        # IMPROVEMENT SUGGESTIONS
        # --------------------------------------------------
        st.subheader("💡 Recommended Next Steps")

        if missing_keywords:

            st.write(
                "1. Compare the missing terms with your real "
                "skills, projects and experience."
            )

            st.write(
                "2. Where appropriate, describe relevant projects "
                "using terminology from the job advertisement."
            )

            st.write(
                "3. If you genuinely have the required experience, "
                "make your contribution and results clearer."
            )

        else:

            st.write(
                "1. Review your CV for clarity and spelling."
            )

            st.write(
                "2. Ensure your projects and experience demonstrate "
                "the requirements of the target role."
            )

            st.write(
                "3. Keep your CV accurate, readable and relevant."
            )

        # --------------------------------------------------
        # DOWNLOADABLE REPORT
        # --------------------------------------------------
        report_data = {
            "Metric": [
                "Keyword Match Score (%)",
                "Matched Keyword Count",
                "Potential Missing Keyword Count",
                "Matched Keywords",
                "Potential Missing Keywords"
            ],
            "Result": [
                round(match_score, 1),
                len(matched_keywords),
                len(missing_keywords),
                ", ".join(matched_keywords),
                ", ".join(missing_keywords)
            ]
        }

        report_df = pd.DataFrame(report_data)

        st.download_button(
            label="⬇️ Download CV Analysis Report",
            data=report_df.to_csv(index=False).encode("utf-8"),
            file_name="skillbridge_cv_analysis.csv",
            mime="text/csv"
        )

        st.caption(
            "Your analysis is based on the CV text and job description "
            "provided in this session. This prototype does not verify "
            "qualifications or guarantee recruitment outcomes."
        )


